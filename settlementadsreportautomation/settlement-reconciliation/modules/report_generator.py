"""Excel report generation for reconciliation outputs."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from modules.ai_insights import build_executive_summary
from modules.dashboard import build_campaign_summary
from modules.reconciler import ReconciliationResult
from utils.constants import STATUS_DUPLICATE, STATUS_MATCH


class ReportGenerator:
    """Create downloadable Excel workbooks from reconciliation results."""

    REPORT_TYPES = {
        "Matched_Report.xlsx": "matched",
        "Unmatched_Report.xlsx": "unmatched",
        "Duplicate_Report.xlsx": "duplicate",
        "Variance_Report.xlsx": "variance",
        "Summary_Report.xlsx": "summary",
        "Campaign_Summary_Report.xlsx": "campaign_summary",
        "Executive_Summary_Report.xlsx": "executive_summary",
    }
    EXCEL_MAX_ROWS = 1_048_576
    EXCEL_MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1
    WIDTH_SAMPLE_ROWS = 100

    def generate_report(
        self,
        filename: str,
        result: ReconciliationResult,
        run_id: str,
    ) -> bytes:
        """Generate one workbook on demand to limit peak memory for large runs."""
        if filename not in self.REPORT_TYPES:
            raise ValueError(f"Unsupported report filename: {filename}.")
        records = result.records
        report_type = self.REPORT_TYPES[filename]
        if report_type == "matched":
            selection_mask = records["Reconciliation_Status"].eq(STATUS_MATCH)
        elif report_type == "unmatched":
            selection_mask = records["Reconciliation_Status"].ne(STATUS_MATCH)
        elif report_type == "duplicate":
            selection_mask = records["Reconciliation_Status"].eq(STATUS_DUPLICATE)
        elif report_type == "variance":
            selection_mask = (
                records["Amount_Variance"].fillna(0).abs().gt(0)
                | records["Reconciliation_Status"].ne(STATUS_MATCH)
            )
        elif report_type == "campaign_summary":
            frame = build_campaign_summary(records)
            return self._workbook(frame, None, title="Campaign Summary")
        elif report_type == "summary":
            frame = pd.DataFrame(
                [
                    {"Metric": key, "Value": value}
                    for key, value in result.metrics.items()
                ]
            )
            frame["Execution_Timestamp_UTC"] = result.execution_timestamp.isoformat()
            return self._workbook(frame, result.summary, title="Summary")
        else:
            executive = build_executive_summary(result)
            frame = pd.DataFrame(
                [
                    {"Section": "Executive Summary", "Metric": "Headline", "Value": executive.headline},
                    {"Section": "Executive Summary", "Metric": "Narrative", "Value": executive.narrative},
                    *[
                        {"Section": "KPI", "Metric": key, "Value": value}
                        for key, value in result.metrics.items()
                    ],
                ]
            )
            frame["Run_ID"] = run_id
            frame["Execution_Timestamp_UTC"] = result.execution_timestamp.isoformat()
            return self._workbook(
                frame,
                result.summary,
                extra_sheets=[
                    ("Priority Insights", executive.insights),
                    ("Campaign Summary", executive.campaign_summary),
                ],
                title="Executive Summary",
            )
        return self._workbook(records, None, selection_mask=selection_mask)

    def generate_outcome_report(
        self,
        result: ReconciliationResult,
        run_id: str,
    ) -> bytes:
        """Generate the master workbook and dedicated PC, ANPL, and SARAL tabs."""
        metrics = pd.DataFrame(
            [
                {"Metric": key, "Value": value}
                for key, value in result.metrics.items()
            ]
        )
        metrics["Run_ID"] = run_id
        metrics["Execution_Timestamp_UTC"] = result.execution_timestamp.isoformat()
        campaign_frames = self._campaign_sheets(result.records)
        calculation_notes = pd.DataFrame(
            [
                {"Calculation": "Settlement source", "Method": "Ads worksheet only"},
                {
                    "Calculation": "Settlement value",
                    "Method": "Use the Ads settlement value; calculate uncached row-level "
                    "SUM(G:K) formulas from their component cells. For recognized SPWT rows, "
                    "compare the absolute value to wallet gross; retain the original signed "
                    "source amount in Ads_Settlement_Amount_Raw. Use the date encoded in "
                    "dated SPWT transaction IDs for transaction-date comparison; retain Ads "
                    "Payment Date as a source field.",
                },
                {
                    "Calculation": "Wallet amount",
                    "Method": "Gross Amount (Incl. of Tax)",
                },
                {
                    "Calculation": "Amount variance",
                    "Method": "Settlement amount minus wallet amount; positive means "
                    "settlement is higher.",
                },
                {
                    "Calculation": "Paired transaction variance",
                    "Method": "Amount difference totals include only transaction IDs present "
                    "in both sources; unmatched source amounts are reported separately.",
                },
                {
                    "Calculation": "Amount match tolerance",
                    "Method": result.metrics.get("amount_tolerance", "Not recorded"),
                },
                {
                    "Calculation": "Campaign category",
                    "Method": "SPWT transaction ID suffix; original Ads and wallet fields "
                    "remain in the master sheet.",
                },
            ]
        )
        return self._workbook(
            result.records,
            result.summary,
            extra_sheets=[
                ("Campaign Summary", campaign_frames[0][1]),
                ("Run Metrics", metrics),
                *campaign_frames[1:],
                ("Calculation Notes", calculation_notes),
            ],
            title="Reconciliation Outcome",
        )

    @staticmethod
    def _campaign_sheets(records: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
        """Split the master outcome by normalized campaign category."""
        if "Campaign_ID" in records and "Transaction_ID" in records:
            summary = build_campaign_summary(records)
            categories = records["Campaign_ID"].astype("string").str.upper()
            campaign_frames = [
                (
                    category,
                    records.loc[categories.eq(category).fillna(False)],
                )
                for category in ("PC", "ANPL", "SARAL")
            ]
            other_mask = categories.isna() | ~categories.isin(("PC", "ANPL", "SARAL"))
            campaign_frames.append(("Other Campaigns", records.loc[other_mask]))
        else:
            summary = pd.DataFrame()
            empty_records = records.iloc[0:0]
            campaign_frames = [
                (category, empty_records)
                for category in ("PC", "ANPL", "SARAL", "Other Campaigns")
            ]
        return [("Campaign Summary", summary), *campaign_frames]

    @staticmethod
    def _workbook(
        frame: pd.DataFrame,
        status_summary: pd.DataFrame | None,
        selection_mask: pd.Series | None = None,
        extra_sheets: list[tuple[str, pd.DataFrame]] | None = None,
        title: str = "Report",
    ) -> bytes:
        """Stream workbook rows and split reports that exceed Excel's row limit."""
        output = BytesIO()
        workbook = Workbook(write_only=True)
        ReportGenerator._append_frame(workbook, frame, title, selection_mask)
        if status_summary is not None:
            ReportGenerator._append_frame(workbook, status_summary, "Status Summary")
        for sheet_title, sheet_frame in extra_sheets or []:
            ReportGenerator._append_frame(workbook, sheet_frame, sheet_title)
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _excel_value(value: Any) -> Any:
        """Convert Pandas and NumPy values to safe, Excel-compatible cell values."""
        if pd.isna(value):
            return None
        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, str):
            if len(value) > 32_767:
                raise ValueError("A report cell exceeds Excel's 32,767-character cell limit.")
            if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
                raise ValueError("A report cell contains characters Excel cannot store.")
            if value.lstrip().startswith(("=", "+", "-", "@")):
                return f"'{value}"
        return value

    @classmethod
    def _append_frame(
        cls,
        workbook: Workbook,
        frame: pd.DataFrame,
        title: str,
        selection_mask: pd.Series | None = None,
    ) -> None:
        """Append one table to one or more write-only worksheets."""
        headers = [str(column) for column in frame.columns]
        if not headers:
            headers = ["No data"]
            frame = pd.DataFrame({"No data": []})

        def create_sheet(part: int) -> tuple[Any, list[int]]:
            sheet_name = title if part == 0 else f"{title[:24]} {part + 1}"
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.freeze_panes = "A2"
            widths = [max(10, min(len(header) + 2, 48)) for header in headers]
            worksheet.append(headers)
            return worksheet, widths

        worksheet, widths = create_sheet(0)
        part = 0
        rows_in_sheet = 0
        if selection_mask is None:
            selected = frame.itertuples(index=False, name=None)
        else:
            mask_values = selection_mask.to_numpy(dtype=bool, na_value=False)
            selected = (
                row
                for row, include in zip(
                    frame.itertuples(index=False, name=None), mask_values, strict=True
                )
                if include
            )

        for row in selected:
            if rows_in_sheet == cls.EXCEL_MAX_DATA_ROWS:
                worksheet.auto_filter.ref = (
                    f"A1:{get_column_letter(len(headers))}{rows_in_sheet + 1}"
                )
                for column_index, width in enumerate(widths, start=1):
                    worksheet.column_dimensions[get_column_letter(column_index)].width = width
                part += 1
                worksheet, widths = create_sheet(part)
                rows_in_sheet = 0
            values = [cls._excel_value(value) for value in row]
            worksheet.append(values)
            rows_in_sheet += 1
            if rows_in_sheet <= cls.WIDTH_SAMPLE_ROWS:
                for column_index, value in enumerate(values):
                    if value is not None:
                        widths[column_index] = min(
                            max(widths[column_index], len(str(value)) + 2),
                            48,
                        )
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{rows_in_sheet + 1}"
        )
        for column_index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
