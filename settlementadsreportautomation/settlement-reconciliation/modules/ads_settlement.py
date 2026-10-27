"""Normalize Flipkart Ads settlement worksheets for reconciliation."""

from __future__ import annotations

import re
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from modules.validator import FileValidationError

_SPWT_ID = re.compile(r"SPWT(?P<sequence>\d+)(?P<ending>[A-Za-z]+)", re.IGNORECASE)
_SARAL_ADS_ID = re.compile(r"ADS_GMV_SETTLE_(?P<sequence>\d+)", re.IGNORECASE)
_SPWT_CAMPAIGNS = frozenset({"PC", "ANPL", "SARAL"})
_SUM_G_TO_K = re.compile(
    r"^=\s*SUM\(\s*\$?G\$?(?P<start>\d+)\s*:\s*\$?K\$?(?P<end>\d+)\s*\)\s*$",
    re.IGNORECASE,
)


def _canonical_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def campaign_id_from_transaction(transaction_id: object) -> str | None:
    """Return the campaign encoded in a supported settlement transaction ID."""
    if transaction_id is None or pd.isna(transaction_id):
        return None
    match = _SPWT_ID.fullmatch(str(transaction_id).strip())
    if match is not None:
        campaign_id = match.group("ending").upper()
        return campaign_id if campaign_id in _SPWT_CAMPAIGNS else None
    if _SARAL_ADS_ID.fullmatch(str(transaction_id).strip()):
        return "SARAL"
    return None


def transaction_date_from_transaction_id(transaction_id: object) -> pd.Timestamp | None:
    """Read the YYMMDD transaction date from a supported ID when present."""
    if transaction_id is None or pd.isna(transaction_id):
        return None
    transaction_value = str(transaction_id).strip()
    match = _SPWT_ID.fullmatch(transaction_value)
    if match is not None:
        sequence = match.group("sequence")
    else:
        ads_match = _SARAL_ADS_ID.fullmatch(transaction_value)
        if ads_match is None:
            return None
        sequence = ads_match.group("sequence")
    if len(sequence) < 6:
        return None
    parsed = pd.to_datetime(
        sequence[:6],
        format="%y%m%d",
        errors="coerce",
    )
    return None if pd.isna(parsed) else pd.Timestamp(parsed).normalize()


def _calculated_settlement_value(
    excel_row: int,
    amount_cell: object,
    filename: str,
    component_cells: tuple[object, ...] = (),
    cached_amount: object = None,
    cached_components: tuple[object, ...] | None = None,
) -> float:
    """Read the value or calculate an uncached Ads row-level SUM(G:K) formula."""
    value = amount_cell.value
    if getattr(amount_cell, "data_type", None) != "f":
        try:
            amount = float(str(value).replace(",", "").strip())
        except (TypeError, ValueError) as exc:
            raise FileValidationError(
                f"{filename}: non-numeric settlement value on Ads row {excel_row}."
            ) from exc
    else:
        formula_match = _SUM_G_TO_K.fullmatch(str(value).strip())
        if (
            formula_match is None
            or int(formula_match.group("start")) != excel_row
            or int(formula_match.group("end")) != excel_row
        ):
            raise FileValidationError(
                f"{filename}: Ads row {excel_row} has an uncached settlement formula "
                "that is not a row-level SUM(G:K); recalculate the workbook in Excel "
                "and upload it again."
            )
        amount = 0.0
        unresolved_component = False
        for index, component in enumerate(component_cells):
            cached_component = (
                cached_components[index]
                if cached_components is not None and index < len(cached_components)
                else None
            )
            if component.data_type == "f" and cached_component is None:
                unresolved_component = True
                continue
            if component.data_type == "e" or (
                isinstance(cached_component, str) and cached_component.startswith("#")
            ):
                raise FileValidationError(
                    f"{filename}: Ads row {excel_row} has an Excel error in "
                    f"{component.column_letter}, used by its settlement formula."
                )
            component_value = (
                cached_component if cached_component is not None else component.value
            )
            if isinstance(component_value, (int, float, np.number)) and not isinstance(
                component_value, bool
            ):
                amount += float(component_value)
        if unresolved_component:
            if isinstance(cached_amount, (int, float, np.number)) and not isinstance(
                cached_amount, bool
            ):
                amount = float(cached_amount)
            else:
                raise FileValidationError(
                    f"{filename}: Ads row {excel_row} has an uncached formula in "
                    "columns G:K and no usable settlement result; recalculate the "
                    "workbook in Excel and upload it again."
                )

    if not np.isfinite(amount):
        raise FileValidationError(
            f"{filename}: non-finite settlement value on Ads row {excel_row}."
        )
    return amount


def load_ads_settlement(file_obj: BinaryIO, filename: str) -> pd.DataFrame:
    """Load the Ads sheet and map its SPWT ending into the reconciliation campaign key."""
    if Path(filename).suffix.lower() != ".xlsx":
        raise FileValidationError("The Ads worksheet format requires an .xlsx settlement report.")

    try:
        file_obj.seek(0)
        workbook = load_workbook(file_obj, read_only=True, data_only=False)
    except (BadZipFile, EOFError, OSError, ValueError, KeyError) as exc:
        raise FileValidationError(f"{filename}: unable to read the Ads settlement workbook.") from exc

    try:
        if "Ads" not in workbook.sheetnames:
            raise FileValidationError(
                f"{filename}: required columns were not found and the workbook has no 'Ads' sheet."
            )
        worksheet = workbook["Ads"]
        rows = worksheet.iter_rows(values_only=True)
        header_row: tuple[object, ...] | None = None
        header_row_number = 0
        required = {
            "paymentdate": "Payment Date",
            "campaigntransactionid": "Campaign / Transaction ID",
            _canonical_header("Settlement Value (Rs.) = SUM(G:K)"): (
                "Settlement Value (Rs.) = SUM(G:K)"
            ),
        }
        for header_row_number, candidate in enumerate(rows, start=1):
            if header_row_number > 20:
                break
            normalized = {
                _canonical_header(value): index for index, value in enumerate(candidate)
            }
            if all(name in normalized for name in required):
                header_row = candidate
                break
        if header_row is None:
            raise FileValidationError(
                f"{filename}: the 'Ads' sheet is missing Payment Date, "
                "Campaign / Transaction ID, or Settlement Value columns."
            )

        header_indexes = {
            _canonical_header(value): index for index, value in enumerate(header_row)
        }
        date_index = header_indexes["paymentdate"]
        transaction_index = header_indexes["campaigntransactionid"]
        amount_index = header_indexes[
            _canonical_header("Settlement Value (Rs.) = SUM(G:K)")
        ]
        source_headers: list[str] = []
        used_headers: set[str] = set()
        for index, value in enumerate(header_row, start=1):
            label = str(value).strip() if value is not None else ""
            if not label:
                label = f"Column_{index}"
            source_name = f"Ads_{label}"
            duplicate_number = 2
            unique_name = source_name
            while unique_name in used_headers:
                unique_name = f"{source_name}_{duplicate_number}"
                duplicate_number += 1
            source_headers.append(unique_name)
            used_headers.add(unique_name)

        records: list[dict[str, object]] = []
        formula_rows: dict[
            int, tuple[int, tuple[int, ...], object, tuple[object, ...]]
        ] = {}
        blank_amount_rows = 0
        incomplete_rows = 0
        for excel_row, cells in enumerate(
            worksheet.iter_rows(min_row=header_row_number + 1),
            start=header_row_number + 1,
        ):
            values = tuple(cell.value for cell in cells)
            if not any(value is not None for value in values):
                continue
            date_value = values[date_index] if date_index < len(values) else None
            amount_cell = cells[amount_index] if amount_index < len(cells) else None
            amount_value = amount_cell.value if amount_cell is not None else None
            if date_value in (None, "") and amount_value in (None, ""):
                incomplete_rows += 1
                continue
            transaction_value = values[transaction_index] if transaction_index < len(values) else None
            transaction_id = (
                str(transaction_value).strip()
                if transaction_value is not None and str(transaction_value).strip()
                else pd.NA
            )
            campaign_id = campaign_id_from_transaction(transaction_id)
            if campaign_id is None:
                campaign_id = transaction_id
            if pd.isna(transaction_id):
                campaign_id = pd.NA

            if amount_value in (None, ""):
                amount = 0.0
                blank_amount_rows += 1
            elif amount_cell.data_type == "f":
                amount = 0.0
            else:
                amount = _calculated_settlement_value(
                    excel_row, amount_cell, filename
                )
            record: dict[str, object] = {
                "Date": date_value,
                "Campaign_ID": campaign_id,
                "Transaction_ID": transaction_id,
                "Customer_ID": pd.NA,
                "Amount": amount,
                "Settlement_Status": "SETTLED",
            }
            for index, source_name in enumerate(source_headers):
                source_value = values[index] if index < len(values) else None
                record[source_name] = source_value
            formula_columns = tuple(
                index for index, cell in enumerate(cells) if cell.data_type == "f"
            )
            if formula_columns:
                amount_components: tuple[object, ...] = ()
                if amount_cell.data_type == "f":
                    amount_components = next(
                        worksheet.iter_rows(
                            min_row=excel_row,
                            max_row=excel_row,
                            min_col=7,
                            max_col=11,
                        )
                    )
                formula_rows[excel_row] = (
                    len(records),
                    formula_columns,
                    amount_cell,
                    amount_components,
                )
            records.append(record)
    finally:
        workbook.close()

    if formula_rows:
        try:
            file_obj.seek(0)
            cached_workbook = load_workbook(file_obj, read_only=True, data_only=True)
        except (BadZipFile, EOFError, OSError, ValueError, KeyError) as exc:
            raise FileValidationError(
                f"{filename}: unable to read cached Ads formula results."
            ) from exc
        try:
            cached_worksheet = cached_workbook["Ads"]
            for excel_row, cells in enumerate(
                cached_worksheet.iter_rows(min_row=header_row_number + 1),
                start=header_row_number + 1,
            ):
                formula_details = formula_rows.get(excel_row)
                if formula_details is None:
                    continue
                (
                    record_index,
                    formula_columns,
                    amount_cell,
                    amount_components,
                ) = formula_details
                record = records[record_index]
                cached_values = tuple(cell.value for cell in cells)
                for index in formula_columns:
                    cached_value = cached_values[index] if index < len(cached_values) else None
                    if cached_value is None:
                        continue
                    record[source_headers[index]] = cached_value
                    if index == date_index:
                        record["Date"] = cached_value
                    elif index == transaction_index:
                        transaction_id = str(cached_value).strip()
                        record["Transaction_ID"] = transaction_id or pd.NA
                        record["Campaign_ID"] = (
                            campaign_id_from_transaction(transaction_id) or transaction_id
                        )
                    elif index == amount_index:
                        record["Amount"] = cached_value

                if amount_cell.data_type == "f":
                    cached_amount = (
                        cached_values[amount_index]
                        if amount_index < len(cached_values)
                        else None
                    )
                    cached_components = next(
                        cached_worksheet.iter_rows(
                            min_row=excel_row,
                            max_row=excel_row,
                            min_col=7,
                            max_col=11,
                            values_only=True,
                        )
                    )
                    record["Amount"] = _calculated_settlement_value(
                        excel_row,
                        amount_cell,
                        filename,
                        amount_components,
                        cached_amount,
                        cached_components,
                    )
                    record[source_headers[amount_index]] = record["Amount"]
        finally:
            cached_workbook.close()

    if not records:
        raise FileValidationError(f"{filename}: the 'Ads' sheet contains no transaction rows.")

    frame = pd.DataFrame.from_records(records)
    payment_dates = pd.to_datetime(frame["Date"], errors="coerce").dt.normalize()
    transaction_dates = frame["Transaction_ID"].map(
        transaction_date_from_transaction_id
    )
    frame["Date"] = payment_dates
    encoded_date_mask = transaction_dates.notna()
    frame.loc[encoded_date_mask, "Date"] = transaction_dates.loc[encoded_date_mask]
    invalid_dates = int(frame["Date"].isna().sum())
    if invalid_dates:
        raise FileValidationError(
            f"{filename}: {invalid_dates} Ads row(s) have invalid or blank Payment Date values."
        )
    frame["Amount"] = pd.to_numeric(frame["Amount"], errors="raise").astype("float64")
    frame["Ads_Settlement_Amount_Raw"] = frame["Amount"]
    spwt_campaigns = frame["Campaign_ID"].isin(("PC", "ANPL", "SARAL"))
    frame.loc[spwt_campaigns, "Amount"] = frame.loc[spwt_campaigns, "Amount"].abs()
    frame["Campaign_ID"] = frame["Campaign_ID"].astype("string")
    frame["Transaction_ID"] = frame["Transaction_ID"].astype("string")
    frame["Customer_ID"] = frame["Customer_ID"].astype("string")
    import_warnings: list[str] = []
    if blank_amount_rows:
        import_warnings.append(
            f"{blank_amount_rows:,} Ads row(s) had a blank settlement value and were treated as 0.00."
        )
    if spwt_campaigns.any():
        import_warnings.append(
            "Recognized campaign Ads amounts are compared by absolute value; original signed "
            "values are retained in Ads_Settlement_Amount_Raw. The transaction date is "
            "read from the transaction ID when available; the original Payment Date is retained."
        )
    if incomplete_rows:
        import_warnings.append(
            f"{incomplete_rows:,} Ads row(s) with both blank Payment Date and settlement value were excluded."
        )
    frame.attrs["import_warnings"] = import_warnings
    return frame
