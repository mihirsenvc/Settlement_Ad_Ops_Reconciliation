"""Tests for streamed Excel output and worksheet splitting."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from modules.reconciler import ReconciliationResult
from modules.report_generator import ReportGenerator
from utils.constants import STATUS_MATCH


def test_report_streaming_splits_rows_and_neutralizes_formulas(monkeypatch) -> None:
    result = ReconciliationResult(
        records=pd.DataFrame(
            {
                "Transaction_ID": ["T1", "T2", "T3"],
                "Reconciliation_Status": [STATUS_MATCH] * 3,
                "Settlement_Amount": [1.0, 2.0, 3.0],
                "Wallet_Amount": [1.0, 2.0, 3.0],
                "Amount_Variance": [0.0, 0.0, 0.0],
                "Formula_Text": ["=1+1", "safe", "safe"],
            }
        ),
        summary=pd.DataFrame(),
        metrics={},
        execution_timestamp=datetime.now(timezone.utc),
    )
    monkeypatch.setattr(ReportGenerator, "EXCEL_MAX_DATA_ROWS", 2)

    payload = ReportGenerator().generate_report("Matched_Report.xlsx", result, "test-run")
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)

    assert workbook.sheetnames == ["Report", "Report 2"]
    assert workbook["Report"]["F2"].value == "'=1+1"
    assert workbook["Report 2"]["A2"].value == "T3"
    workbook.close()


def test_report_rejects_cells_exceeding_excel_limit() -> None:
    with pytest.raises(ValueError, match="Excel's 32,767-character"):
        ReportGenerator._excel_value("x" * 32_768)


@pytest.mark.parametrize("value", [" =1+1", "\t@SUM(A1:A2)", "\r-2+3"])
def test_report_neutralizes_formulas_after_leading_whitespace(value: str) -> None:
    assert ReportGenerator._excel_value(value) == f"'{value}"


def test_report_rejects_xml_control_characters() -> None:
    with pytest.raises(ValueError, match="characters Excel cannot store"):
        ReportGenerator._excel_value("unsafe\x01value")


def test_new_executive_and_campaign_reports_are_generated() -> None:
    records = pd.DataFrame(
        {
            "Date": pd.to_datetime(["2026-10-01"]),
            "Campaign_ID": ["C1"],
            "Campaign_Name": ["Campaign One"],
            "Business_Unit": ["Consumer"],
            "Transaction_ID": ["T1"],
            "Customer_ID": ["U1"],
            "Settlement_Amount": [10.0],
            "Wallet_Amount": [10.0],
            "Amount_Variance": [0.0],
            "Reconciliation_Status": [STATUS_MATCH],
        }
    )
    result = ReconciliationResult(
        records=records,
        summary=pd.DataFrame(
            {
                "Reconciliation_Status": [STATUS_MATCH],
                "Records": [1],
            }
        ),
        metrics={
            "total_records": 1,
            "matched_records": 1,
            "unmatched_records": 0,
            "duplicate_records": 0,
            "match_percentage": 100.0,
            "variance_amount": 0.0,
            "variance_percentage": 0.0,
            "transaction_quality_score": 100.0,
        },
        execution_timestamp=datetime.now(timezone.utc),
    )
    generator = ReportGenerator()

    campaign_bytes = generator.generate_report(
        "Campaign_Summary_Report.xlsx", result, "test-run"
    )
    summary_bytes = generator.generate_report("Summary_Report.xlsx", result, "test-run")
    executive_bytes = generator.generate_report(
        "Executive_Summary_Report.xlsx", result, "test-run"
    )
    campaign_book = load_workbook(BytesIO(campaign_bytes), read_only=True)
    summary_book = load_workbook(BytesIO(summary_bytes), read_only=True)
    executive_book = load_workbook(BytesIO(executive_bytes), read_only=True)

    assert campaign_book.sheetnames == ["Campaign Summary"]
    assert summary_book.sheetnames == ["Summary", "Status Summary"]
    assert "Executive Summary" in executive_book.sheetnames
    assert "Priority Insights" in executive_book.sheetnames
    assert "Campaign Summary" in executive_book.sheetnames
    campaign_book.close()
    summary_book.close()
    executive_book.close()


def test_complete_outcome_report_contains_transactions_summary_and_run_metrics() -> None:
    result = ReconciliationResult(
        records=pd.DataFrame(
            {
                "Transaction_ID": ["T1"],
                "Reconciliation_Status": [STATUS_MATCH],
                "Settlement_Amount": [10.0],
                "Wallet_Amount": [10.0],
                "Amount_Variance": [0.0],
            }
        ),
        summary=pd.DataFrame(
            {
                "Reconciliation_Status": [STATUS_MATCH],
                "Records": [1],
            }
        ),
        metrics={"total_records": 1, "matched_records": 1},
        execution_timestamp=datetime.now(timezone.utc),
    )

    payload = ReportGenerator().generate_outcome_report(result, "test-run")
    workbook = load_workbook(BytesIO(payload), read_only=True)

    assert workbook.sheetnames == [
        "Reconciliation Outcome",
        "Status Summary",
        "Campaign Summary",
        "Run Metrics",
        "PC",
        "ANPL",
        "SARAL",
        "Other Campaigns",
        "Calculation Notes",
    ]
    assert workbook["Reconciliation Outcome"]["A2"].value == "T1"
    assert workbook["Status Summary"]["A2"].value == STATUS_MATCH
    assert workbook["Run Metrics"]["A2"].value == "total_records"
    assert workbook["Run Metrics"]["C2"].value == "test-run"
    workbook.close()


def test_complete_outcome_report_separates_pc_anpl_and_saral_and_keeps_source_fields() -> None:
    records = pd.DataFrame(
        {
            "Transaction_ID": ["SPWT1PC", "SPWT2ANPL", "SPWT3SARAL"],
            "Campaign_ID": ["PC", "ANPL", "SARAL"],
            "Reconciliation_Status": [STATUS_MATCH] * 3,
            "Settlement_Amount": [10.0, 20.0, 30.0],
            "Wallet_Amount": [10.0, 20.0, 30.0],
            "Amount_Variance": [0.0, 0.0, 0.0],
            "Ads_Type": ["pc charge", "anpl charge", "saral charge"],
            "Wallet_Operation": ["PC", "ANPL", "SARAL"],
        }
    )
    result = ReconciliationResult(
        records=records,
        summary=pd.DataFrame(
            {
                "Reconciliation_Status": [STATUS_MATCH],
                "Records": [3],
                "Settlement_Amount": [60.0],
                "Wallet_Amount": [60.0],
                "Variance_Amount": [0.0],
            }
        ),
        metrics={"total_records": 3, "amount_tolerance": 0.01},
        execution_timestamp=datetime.now(timezone.utc),
    )

    workbook = load_workbook(
        BytesIO(ReportGenerator().generate_outcome_report(result, "test-run")),
        read_only=True,
    )

    assert workbook["PC"]["A2"].value == "SPWT1PC"
    assert workbook["ANPL"]["A2"].value == "SPWT2ANPL"
    assert workbook["SARAL"]["A2"].value == "SPWT3SARAL"
    master_headers = [cell.value for cell in workbook["Reconciliation Outcome"][1]]
    master_row = [cell.value for cell in workbook["Reconciliation Outcome"][2]]
    assert "Ads_Type" in master_headers
    assert "Wallet_Operation" in master_headers
    assert master_row[master_headers.index("Ads_Type")] == "pc charge"
    assert workbook["Calculation Notes"]["A2"].value == "Settlement source"
    workbook.close()
