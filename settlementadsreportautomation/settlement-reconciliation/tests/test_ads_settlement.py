"""Tests for Flipkart Ads settlement worksheet normalization."""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook

from app import _read_upload
from modules.ads_settlement import (
    campaign_id_from_transaction,
    load_ads_settlement,
    transaction_date_from_transaction_id,
)
from modules.validator import FileValidationError


def _ads_workbook(rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ads"
    worksheet.append(["Settlement report"])
    worksheet.append(
        [
            "NEFT ID",
            "Payment Date",
            "Settlement Value (Rs.) = SUM(G:K)",
            None,
            "Type",
            "Campaign / Transaction ID",
        ]
    )
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_load_ads_settlement_normalizes_spwt_and_seller_campaign_rows() -> None:
    source = _ads_workbook(
        [
            ["NEFT-1", "2026-05-27", -125.5, None, "topup", "SPWT260403230546862429624ANPL"],
            ["NEFT-1", "2026-05-27", -30, None, "redeem", "CMP-SELLER-1"],
        ]
    )

    frame = load_ads_settlement(source, "settlement.xlsx")

    assert list(frame.columns[:6]) == [
        "Date",
        "Campaign_ID",
        "Transaction_ID",
        "Customer_ID",
        "Amount",
        "Settlement_Status",
    ]
    assert frame.loc[0, "Campaign_ID"] == "ANPL"
    assert frame.loc[0, "Transaction_ID"] == "SPWT260403230546862429624ANPL"
    assert frame.loc[0, "Amount"] == 125.5
    assert frame.loc[0, "Ads_Settlement_Amount_Raw"] == -125.5
    assert frame.loc[0, "Ads_Settlement Value (Rs.) = SUM(G:K)"] == -125.5
    assert frame.loc[1, "Campaign_ID"] == "CMP-SELLER-1"
    assert frame.loc[0, "Ads_Type"] == "topup"
    assert frame.loc[0, "Date"] == pd.Timestamp("2026-04-03")
    assert frame.loc[1, "Date"] == pd.Timestamp("2026-05-27")
    assert frame.loc[0, "Ads_Payment Date"] == "2026-05-27"


def test_load_ads_settlement_maps_ads_gmv_settle_ids_to_saral() -> None:
    transaction_id = "ADS_GMV_SETTLE_26090600003893499075"
    source = _ads_workbook(
        [["NEFT-1", "2026-09-10", -125.5, None, "settlement", transaction_id]]
    )

    frame = load_ads_settlement(source, "settlement.xlsx")

    assert frame.loc[0, "Campaign_ID"] == "SARAL"
    assert frame.loc[0, "Transaction_ID"] == transaction_id
    assert frame.loc[0, "Date"] == pd.Timestamp("2026-09-06")
    assert frame.loc[0, "Amount"] == 125.5
    assert frame.loc[0, "Ads_Settlement_Amount_Raw"] == -125.5


@pytest.mark.parametrize(
    ("transaction_id", "expected"),
    [
        ("SPWT123PC", "PC"),
        ("spwt123anpl", "ANPL"),
        ("SPWT123SARAL", "SARAL"),
        ("ADS_GMV_SETTLE_26090600003893499075", "SARAL"),
        ("OTHER123PC", None),
        ("SPWT123UNKNOWN", None),
        ("ADS_GMV_SETTLE_invalid", None),
    ],
)
def test_campaign_id_requires_spwt_prefix_and_supported_suffix(
    transaction_id: str, expected: str | None
) -> None:
    assert campaign_id_from_transaction(transaction_id) == expected


def test_spwt_transaction_date_uses_encoded_yymmdd_and_falls_back_for_short_ids() -> None:
    assert transaction_date_from_transaction_id("SPWT251217110353667749627ANPL") == pd.Timestamp(
        "2025-12-17"
    )
    assert transaction_date_from_transaction_id(
        "ADS_GMV_SETTLE_26090600003893499075"
    ) == pd.Timestamp("2026-09-06")
    assert transaction_date_from_transaction_id("SPWT123PC") is None


def test_load_ads_settlement_treats_blank_amount_as_zero_and_warns() -> None:
    source = _ads_workbook(
        [["NEFT-1", "2026-05-27", None, None, "topup", "SPWT123PC"]]
    )

    frame = load_ads_settlement(source, "settlement.xlsx")

    assert frame.loc[0, "Amount"] == 0
    assert frame.attrs["import_warnings"] == [
        "1 Ads row(s) had a blank settlement value and were treated as 0.00.",
        "Recognized campaign Ads amounts are compared by absolute value; original signed "
        "values are retained in Ads_Settlement_Amount_Raw. The transaction date is "
        "read from the transaction ID when available; the original Payment Date is retained.",
    ]


def test_load_ads_settlement_excludes_rows_without_date_or_amount() -> None:
    source = _ads_workbook(
        [
            ["NEFT-1", "2026-05-27", 15, None, "topup", "SPWT123PC"],
            [None, None, None, None, None, "SPWT124ANPL"],
        ]
    )

    frame = load_ads_settlement(source, "settlement.xlsx")

    assert len(frame) == 1
    assert frame.attrs["import_warnings"] == [
        "Recognized campaign Ads amounts are compared by absolute value; original signed "
        "values are retained in Ads_Settlement_Amount_Raw. The transaction date is "
        "read from the transaction ID when available; the original Payment Date is retained.",
        "1 Ads row(s) with both blank Payment Date and settlement value were excluded.",
    ]


def test_upload_reader_uses_ads_sheet_even_when_first_sheet_has_standard_columns() -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Date", "Campaign_ID", "Transaction_ID", "Customer_ID", "Amount", "Settlement_Status"])
    summary.append(["2026-05-27", "WRONG", "WRONG-1", "U1", 999, "SETTLED"])
    ads = workbook.create_sheet("Ads")
    ads.append(["Settlement report"])
    ads.append(
        [
            "NEFT ID",
            "Payment Date",
            "Settlement Value (Rs.) = SUM(G:K)",
            None,
            "Type",
            "Campaign / Transaction ID",
        ]
    )
    ads.append(["NEFT-1", "2026-05-27", 15, None, "topup", "SPWT123PC"])
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    upload = _UploadedWorkbook(source.getvalue(), "settlement.xlsx")

    frame = _read_upload(
        upload,
        "settlement",
        SimpleNamespace(max_xlsx_expanded_size_mb=10),
    )

    assert len(frame) == 1
    assert frame.loc[0, "Campaign_ID"] == "PC"
    assert frame.loc[0, "Amount"] == 15


def test_load_ads_settlement_calculates_uncached_row_sum_formula() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ads"
    worksheet.append(["Settlement report"])
    worksheet.append(
        [
            "NEFT ID",
            "Payment Date",
            "Settlement Value (Rs.) = SUM(G:K)",
            None,
            "Type",
            "Campaign / Transaction ID",
        ]
    )
    worksheet.append(
        ["NEFT-1", "2026-05-27", "=SUM(G3:K3)", None, "topup", "SPWT123ANPL", 10, 20, 30, 40, 50]
    )
    source = BytesIO()
    workbook.save(source)
    source.seek(0)

    frame = load_ads_settlement(source, "settlement.xlsx")

    assert frame.loc[0, "Amount"] == 150
    assert frame.loc[0, "Ads_Settlement Value (Rs.) = SUM(G:K)"] == 150


def test_load_ads_settlement_rejects_uncached_unrecognized_formula() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ads"
    worksheet.append(["Settlement report"])
    worksheet.append(
        [
            "NEFT ID",
            "Payment Date",
            "Settlement Value (Rs.) = SUM(G:K)",
            None,
            "Type",
            "Campaign / Transaction ID",
        ]
    )
    worksheet.append(
        ["NEFT-1", "2026-05-27", "=A3+1", None, "topup", "SPWT123PC"]
    )
    source = BytesIO()
    workbook.save(source)
    source.seek(0)

    with pytest.raises(FileValidationError, match="not a row-level SUM"):
        load_ads_settlement(source, "settlement.xlsx")


def test_load_ads_settlement_rejects_workbooks_without_ads_sheet() -> None:
    workbook = Workbook()
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    with pytest.raises(FileValidationError, match="no 'Ads' sheet"):
        load_ads_settlement(output, "settlement.xlsx")


class _UploadedWorkbook(BytesIO):
    def __init__(self, payload: bytes, filename: str) -> None:
        super().__init__(payload)
        self.name = filename
        self.size = len(payload)
